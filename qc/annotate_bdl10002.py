"""Annotated copies of the BDL_10002 export for curation.

Usage:
    python qc/annotate_bdl10002.py data/original/BDL_10002.xlsx ROY2018_DATA_DIR FULL.xlsx SHORT.xlsx

ROY2018_DATA_DIR is the data/ folder of genenetwork/bxd_gxelongevity_2020.

Only problems that change the numbers going into mortality-rate and mapping analyses
are flagged: which animals count as natural deaths, their ages at death, and their
strain, diet and sex groups. Metadata (RRID, Epoch), shared ear tags and ordinary
outliers are left out on purpose.

FULL.xlsx keeps every original row in its original order and every original cell
unchanged, with seven columns inserted at the far left. SHORT.xlsx holds only the
flagged rows, still in their original order, with the same columns. Both have a
"QC legend" sheet.
"""
import sys
from collections import defaultdict

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

src, roy_dir, out_full, out_short = sys.argv[1:5]
HEADER_ROW = 11  # GeneNetwork puts 10 rows of metadata above the column headers

d = pd.read_excel(src, header=HEADER_ROW - 1)
d["row"] = d.index + HEADER_ROW + 1  # worksheet row of each animal
d["value"] = pd.to_numeric(d.Value, errors="coerce")
d["prefix"] = d.Name.str.extract(r"^([A-Za-z]+)")[0]
d["strain_key"] = d.Strain.str.replace(r"^BXD0*", "BXD", regex=True)

roy_nat = pd.read_csv(f"{roy_dir}/AgingBXD_Longevitydata_Nov2018.natural.deaths.txt", sep="\t")
for c in ("DateOfBirth", "DateOfDeath"):
    roy_nat[c] = pd.to_datetime(roy_nat[c], format="%m.%d.%Y")
roy_alive = pd.read_csv(f"{roy_dir}/AgingBXD_Longevitydata_Nov2018.alive.txt", sep="\t")
roy_excl = pd.read_excel(f"{roy_dir}/AgingBXD_Longevitydata_Nov2018.xlsx", sheet_name="OthersLowOutliers")
for t in (roy_nat, roy_alive, roy_excl):
    t["strain_key"] = t.StrainNameCurrent.astype(str).str.replace(r"^BXD0*", "BXD", regex=True)
key = ["EarTag", "strain_key"]
rkey = ["EarTagNumberCurrent", "strain_key"]

flags = defaultdict(list)  # Name -> list of (priority, code, note)
proposed_value, proposed_longevity = {}, {}


def add(name, priority, code, note):
    flags[name].append((priority, code, note))


def day(x):
    return pd.Timestamp(x).strftime("%-d %b %Y")


nat = d[d.value.notna()].copy()

# DUPLICATE and EXCLUDED_2018: set aside in Nov 2018 as Other or LowOutlier, entered here as natural deaths
ex = nat.merge(roy_excl, left_on=key, right_on=rkey, suffixes=("", "_2018"))
same_day = roy_excl.groupby("DateOfDeath").size()
harvested = set(zip(d[d.prefix == "H"].EarTag, d[d.prefix == "H"].strain_key))
for r in ex.itertuples():
    n = same_day[r.DateOfDeath]
    when = f"died {day(r.DateOfDeath)}" + (f", one of {n} 2018 exclusions that died that day" if n > 1 else "")
    if (r.EarTag, r.strain_key) in harvested:
        add(r.Name, 1, "DUPLICATE",
            f"Same animal as harvested row H{r.EarTag:04d}. Delete this row. 2018 cause: {r.CauseOfDeath}, {when}.")
        proposed_value[r.Name] = "delete row"
    else:
        add(r.Name, 1, "EXCLUDED_2018",
            f"Excluded from natural lifespan in Nov 2018 as \"{r.CauseOfDeath}\" ({when}). Set Value and Longevity to x.")
        proposed_value[r.Name] = proposed_longevity[r.Name] = "x"
excluded = set(ex.Name)

# VALUE_WRONG: Value disagrees with Age
for r in nat[nat.value != nat.Age].itertuples():
    add(r.Name, 1, "VALUE_WRONG",
        f"Value {r.Value} but Age and Longevity are {r.Age}; the Nov 2018 record also gives {r.Age} d.")
    proposed_value[r.Name] = r.Age

# LONGEVITY_BLANK: natural deaths whose Longevity attribute is missing
for r in nat[nat.Longevity.isna() & ~nat.Name.isin(excluded)].itertuples():
    add(r.Name, 1, "LONGEVITY_BLANK",
        f"Longevity is blank; it should equal Value ({r.Value}). Analyses that read Longevity would drop this animal.")
    proposed_longevity[r.Name] = r.Value

# STRAIN_NAME: one strain spelled two ways when case is ignored
for _, g in nat.groupby(nat.Strain.str.lower()):
    if g.Strain.nunique() > 1:
        counts = g.Strain.value_counts()
        for r in g[g.Strain == counts.idxmin()].itertuples():
            add(r.Name, 1, "STRAIN_NAME",
                f"Strain should be {counts.idxmax()} (as for {counts.max()} other animals); otherwise it forms its own group.")

# ALIVE_LATER: died before an age at which it was recorded alive (15 Nov 2018)
m = nat.merge(roy_alive, left_on=key, right_on=rkey, suffixes=("", "_2018"))
for r in m[m.value < m.AgeIfStllAlive].itertuples():
    add(r.Name, 2, "ALIVE_LATER",
        f"Recorded alive at {r.AgeIfStllAlive} d on 15 Nov 2018, but Value is {r.Value}. Check dates.")

# AGE_DIFFERS_2018: age at death differs from the Nov 2018 natural-death record
m = nat.merge(roy_nat, left_on=key, right_on=rkey, suffixes=("", "_2018"))
accepted_2018 = set(m.Name)
for r in m[(m.value != m.AgeAtDeath) & (m.value == m.Age)].itertuples():  # L0852 is VALUE_WRONG
    add(r.Name, 2, "AGE_DIFFERS_2018",
        f"Nov 2018 record: {r.AgeAtDeath} d (born {day(r.DateOfBirth)}, died {day(r.DateOfDeath)}). Check dates.")

# The remaining checks use natural deaths after removing the 2018 exclusions
keep = nat[~nat.Name.isin(excluded)].copy()

# EARLY_DEATH: natural death before 200 d among animals added since 2018
for r in keep[(keep.value < 200) & ~keep.Name.isin(accepted_2018)].itertuples():
    add(r.Name, 2, "EARLY_DEATH",
        f"Died at {int(r.value)} d; added since 2018. Check cause of death against the 2018 rules for LowOutlier.")

# SAME_DAY: three or more animals of one strain and diet dying at the same age (littermates
# share a birth date, so the same day), at least one of them added since 2018
for _, g in keep.groupby(["Strain", "Tx", "value"]):
    if len(g) >= 3 and not set(g.Name) <= accepted_2018:
        for r in g.itertuples():
            others = ", ".join(n for n in g.Name if n != r.Name)
            add(r.Name, 2, "SAME_DAY",
                f"Died at {int(r.value)} d, like {others} of the same strain and diet. "
                "Check whether they died the same day (cage event or cull).")

# STRAIN_MIXUP: outliers whose neighboring ear tags belong to another strain. Found by
# inspecting animals more than 3.5 robust SD from their strain x diet x sex median.
grp = keep.groupby(["Strain", "Tx", "Sex"]).value
keep["med"] = grp.transform("median")
mixups = {"L0101": "the neighboring tags L0097-L0100 are BXD12",
          "L0102": "the neighboring tags L0097-L0100 are BXD12",
          "L0481": "the whole cage L0481-L0484 lived 1,088-1,229 d",
          "L0482": "the whole cage L0481-L0484 lived 1,088-1,229 d",
          "L0483": "the whole cage L0481-L0484 lived 1,088-1,229 d",
          "L0484": "the whole cage L0481-L0484 lived 1,088-1,229 d"}
for r in keep[keep.Name.isin(mixups)].itertuples():
    add(r.Name, 2, "STRAIN_MIXUP",
        f"{int(r.value)} d against a median of {r.med:.0f} d for {r.Strain} {r.Tx} {r.Sex}; {mixups[r.Name]}. "
        "Check the strain in colony records.")

CODES = [  # code, meaning, why it matters downstream
    ("DUPLICATE", "Listed both as a natural death (this row) and as harvested (H row).",
     "Counts a non-natural death as natural."),
    ("EXCLUDED_2018", "Set aside in the Nov 2018 curation as \"Other\" (mostly same-day culls) or \"LowOutlier\".",
     "Adds culls and early non-aging deaths to natural lifespan."),
    ("VALUE_WRONG", "Value disagrees with Age, Longevity and the 2018 record.", "Wrong age at death."),
    ("LONGEVITY_BLANK", "Natural death with Longevity blank; it should equal Value.",
     "Any analysis that reads Longevity drops the animal (68 of 184 NAM animals)."),
    ("STRAIN_NAME", "Strain spelled differently from its siblings.", "Splits one strain into two groups."),
    ("ALIVE_LATER", "Died before an age at which it was recorded alive in Nov 2018.", "Wrong age at death."),
    ("AGE_DIFFERS_2018", "Age at death differs from the Nov 2018 record.", "Possibly wrong age at death."),
    ("EARLY_DEATH", "Death before 200 d, added since 2018.",
     "Early non-aging deaths flatten the Gompertz slope and drive the youngest survivorships."),
    ("SAME_DAY", "Three or more animals of one strain and diet died at the same age, at least one added since 2018.",
     "A cage event puts several deaths in one interval and distorts age-specific mortality."),
    ("STRAIN_MIXUP", "Lifespan far from its strain, next to tags of another strain.",
     "Wrong strain assignment moves the animal between mapping groups."),
]
LEFT_OUT = [
    "RRID and Epoch: blank or mismatched values do not change any age, strain, diet or sex.",
    "Ear tags shared by two animals: accepted in the colony.",
    "Outliers within strain, diet and sex: treated as biology unless they look like a strain mix-up.",
    "Early deaths and same-day deaths already accepted as natural in the Nov 2018 curation.",
    "Same-age deaths of only two animals of a strain and diet.",
    "Three animals alive on 15 Nov 2018 but missing from the export: BXD199 (tag 2773), "
    "BXD16 (tag 2686) and DBA/2J-Gpnmb (tag 2591).",
]

# --- write the workbooks ----------------------------------------------------------------
NEW = ["QC priority", "QC codes", "QC note", "Proposed Value", "Proposed Longevity",
       "Curator decision", "Curator note"]
K = len(NEW)
font = Font(name="Calibri", size=11)
bold = Font(name="Calibri", size=11, bold=True)
qc_fill = PatternFill("solid", fgColor="DCE6F1")      # added QC columns
input_fill = PatternFill("solid", fgColor="FFFF99")   # for the curator to fill in
priority_fill = {1: PatternFill("solid", fgColor="F4B6B6"), 2: PatternFill("solid", fgColor="FCE4B6")}


def annotation(name):
    items = sorted(flags[name])
    codes = list(dict.fromkeys(code for _, code, _ in items))
    return [items[0][0], "; ".join(codes), " | ".join(note for _, _, note in items),
            proposed_value.get(name), proposed_longevity.get(name)]


def format_sheet(ws, header_row, first, last):
    for j, h in enumerate(NEW, start=1):
        c = ws.cell(header_row, j, h)
        c.font = bold
        c.fill = input_fill if j > 5 else qc_fill
    for row in range(first, last + 1):
        p = ws.cell(row, 1).value
        if p in priority_fill:
            ws.cell(row, 1).fill = priority_fill[p]
            ws.cell(row, 1).alignment = Alignment(horizontal="center")
        for j in (6, 7):
            ws.cell(row, j).fill = input_fill
            ws.cell(row, j).font = font
    dv = DataValidation(type="list", formula1='"Accept,Reject,Modify,Pending"', allow_blank=True)
    dv.add(f"F{first}:F{last}")
    ws.add_data_validation(dv)
    for col, width in zip("ABCDEFG", (11, 26, 90, 15, 18, 17, 40)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = ws.cell(first, K + 3)  # keep QC columns, Index and Name in view
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{last}"


def legend(wb, sheet, first, last, what):
    lg = wb.create_sheet("QC legend")
    rng = f"'{sheet}'!$%s${first}:$%s${last}"
    rows = [
        (f"QC annotations for the BDL_10002 export of 22 October 2026 ({what})", None, None, None),
        ("Only problems that change the numbers going into mortality-rate and mapping analyses are flagged.", None, None, None),
        (None, None, None, None),
        ("Priority", "Meaning", "Rows", None),
        (1, "Clear error; a correction is proposed", f"=COUNTIF({rng % ('A', 'A')},1)", None),
        (2, "Check against colony records", f"=COUNTIF({rng % ('A', 'A')},2)", None),
        (None, None, None, None),
        ("Code", "Meaning", "Rows", "Why it matters downstream"),
    ]
    for code, meaning, why in CODES:
        f = f'=SUMPRODUCT(--ISNUMBER(SEARCH("{code};",SUBSTITUTE({rng % ("B", "B")}," ","")&";")))'
        rows.append((code, meaning, f, why))
    rows += [(None, None, None, None), ("Left out on purpose", None, None, None)]
    rows += [(None, s, None, None) for s in LEFT_OUT]
    rows += [
        (None, None, None, None),
        ("Columns added", None, None, None),
        ("QC priority", "1 or 2; the higher priority among the row's flags.", None, None),
        ("QC codes / QC note", "All codes for the row, and one note per code separated by |.", None, None),
        ("Proposed Value / Longevity", "Proposed replacement, where one is proposed. \"delete row\" means remove the row.", None, None),
        ("Curator decision / note", "For the curator (yellow): Accept, Reject, Modify or Pending, and free text.", None, None),
        (None, None, None, None),
        ("Sources: November 2018 data in github.com/genenetwork/bxd_gxelongevity_2020 "
         "(data/AgingBXD_Longevitydata_Nov2018.xlsx and .txt files). Script: qc/annotate_bdl10002.py "
         "in github.com/robwwilliams/bxd_mortality_rates_2026.", None, None, None),
    ]
    heads = {"Priority", "Code", "Left out on purpose", "Columns added"}
    for i, r in enumerate(rows, start=1):
        for j, v in enumerate(r, start=1):
            if v is not None:
                c = lg.cell(i, j, v)
                c.font = bold if (i == 1 or r[0] in heads) else font
                c.alignment = Alignment(wrap_text=j in (2, 4), vertical="top")
    for col, width in zip("ABCD", (28, 80, 8, 60)):
        lg.column_dimensions[col].width = width


# Full length: every original row and cell, seven columns inserted at the left
wb = openpyxl.load_workbook(src)
ws = wb.active
ws.insert_cols(1, K)
ws.cell(1, 1, "QC annotations added 23 October 2026. Columns A-G are new; all other cells are the "
              "original GeneNetwork export, unchanged. See sheet \"QC legend\".").font = bold
by_row = dict(zip(d.Name, d.row))
for name in flags:
    for j, v in enumerate(annotation(name), start=1):
        ws.cell(by_row[name], j, v).font = font
format_sheet(ws, HEADER_ROW, HEADER_ROW + 1, ws.max_row)
legend(wb, ws.title, HEADER_ROW + 1, ws.max_row, "all rows")
wb.save(out_full)

# Short form: only the flagged rows, in their original order
orig = openpyxl.load_workbook(src).active
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Short form"
ws.cell(1, 1, "BDL_10002 short form: only the rows with errors that affect mortality-rate and mapping "
              "analyses, in their original order. Columns A-G are QC annotations; the remaining "
              "columns are the original export. See sheet \"QC legend\".").font = bold
SHORT_HEADER = 3
for j in range(1, orig.max_column + 1):
    ws.cell(SHORT_HEADER, K + j, orig.cell(HEADER_ROW, j).value).font = font
out_row = SHORT_HEADER
for name in sorted(flags, key=lambda n: by_row[n]):
    out_row += 1
    for j, v in enumerate(annotation(name), start=1):
        ws.cell(out_row, j, v).font = font
    for j in range(1, orig.max_column + 1):
        ws.cell(out_row, K + j, orig.cell(by_row[name], j).value).font = font
format_sheet(ws, SHORT_HEADER, SHORT_HEADER + 1, out_row)
legend(wb, ws.title, SHORT_HEADER + 1, out_row, "short form")
wb.save(out_short)

n_by_code = defaultdict(int)
for items in flags.values():
    for c in {code for _, code, _ in items}:
        n_by_code[c] += 1
print("flagged rows:", len(flags))
print("by priority:", pd.Series([min(p for p, _, _ in v) for v in flags.values()]).value_counts().sort_index().to_dict())
print("by code:", dict(sorted(n_by_code.items())))
